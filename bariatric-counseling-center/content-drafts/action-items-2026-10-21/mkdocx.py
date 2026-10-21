"""Markdown draft to a plain Word file for client edits: python mkdocx.py <in.md> <out.docx>"""
import re, sys, docx
from docx.shared import Pt
src, out = sys.argv[1], sys.argv[2]
d = docx.Document()
st = d.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(11)
def add_runs(p, text):
    for part in re.split(r'(\*\*[^*]+\*\*)', text):
        if part.startswith('**') and part.endswith('**'):
            p.add_run(part[2:-2]).bold = True
        elif part:
            p.add_run(part)
for line in open(src, encoding='utf-8').read().split('\n'):
    if not line.strip() or line.strip() == '---':
        continue
    if line.startswith('# '): d.add_heading(line[2:], 0)
    elif line.startswith('## '): d.add_heading(line[3:], 1)
    elif line.startswith('- '): add_runs(d.add_paragraph(style='List Bullet'), line[2:])
    elif re.match(r'^\d+\. ', line): add_runs(d.add_paragraph(style='List Number'), re.sub(r'^\d+\. ', '', line))
    else: add_runs(d.add_paragraph(), line)
d.save(out)
print('wrote', out)
